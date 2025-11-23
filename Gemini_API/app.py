import os
import re
import json
import time
import base64
import tempfile
import asyncio
from io import BytesIO
from datetime import datetime

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv
from PIL import Image
import google.generativeai as genai
import win32com.client as win32

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

load_dotenv()

app = FastAPI()

REQUESTS_PER_MINUTE = 60
semaphore = asyncio.Semaphore(REQUESTS_PER_MINUTE)
last_request_time = 0
MIN_INTERVAL = 60.0 / REQUESTS_PER_MINUTE

app.mount("/static", StaticFiles(directory="static"), name="static")

REPORTS_DIR = "reports"
os.makedirs(REPORTS_DIR, exist_ok=True)

genai.configure(api_key=os.getenv("GEMINI_API_KEY"))


def load_json_file(path: str) -> dict:
    if not os.path.exists(path):
        raise HTTPException(status_code=500, detail=f"Файл {path} не найден.")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def set_document_styles(doc):
    sections = doc.sections
    for section in sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    styles = doc.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(39, 43, 103)
    normal_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    normal_style.paragraph_format.first_line_indent = Cm(1.25)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(0)

    heading1 = styles["Heading 1"]
    heading1.font.name = "Times New Roman"
    heading1.font.size = Pt(18)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(39, 43, 103)
    heading1.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    heading1.paragraph_format.space_after = Pt(24)

    heading2 = styles["Heading 2"]
    heading2.font.name = "Times New Roman"
    heading2.font.size = Pt(16)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(39, 43, 103)
    heading2.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    heading2.paragraph_format.space_before = Pt(24)
    heading2.paragraph_format.space_after = Pt(12)


def add_page_numbers(doc):
    for i, section in enumerate(doc.sections):
        section.different_first_page_header_footer = True

        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph(
        )
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)

        fldSimple = OxmlElement("w:fldSimple")
        fldSimple.set(qn("w:instr"), "PAGE")

        run = paragraph.add_run()
        run._r.append(fldSimple)

        if i == 0:
            section.start_page_number = 2


def add_title_page(doc):
    for _ in range(14):
        doc.add_paragraph()

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(
        "ИССЛЕДОВАНИЕ РЫНКА В СЕГМЕНТЕ, КОТОРОМУ ПРИНАДЛЕЖИТ ОЦЕНИВАЕМАЯ КВАРТИРА")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"

    for _ in range(14):
        doc.add_paragraph()

    for text in [f"Нижний Новгород", str(datetime.today().year)]:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.name = "Times New Roman"


def add_table_of_contents(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run("Содержание")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.bold = True

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldSimple = OxmlElement("w:fldSimple")
    fldSimple.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    run._r.append(fldSimple)


def format_text(doc, text):
    text = text.strip()
    if not text:
        return

    if text.startswith("# "):
        paragraph = doc.add_paragraph(text[2:], style="Heading 1")

    elif text.startswith("## "):
        paragraph = doc.add_paragraph(text[3:], style="Heading 2")

    elif text.startswith("* "):
        paragraph = doc.add_paragraph(style="ListBullet")
        remaining_text = text[2:]

        bold_parts = re.split(r"(\*\*.*?\*\*)", remaining_text)

        for part in bold_parts:
            run = paragraph.add_run()
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]
                run.bold = True
            else:
                run.text = part
            run.font.color.rgb = RGBColor(39, 43, 103)

    else:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        bold_parts = re.split(r"(\*\*.*?\*\*)", text)

        for part in bold_parts:
            run = paragraph.add_run()
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]
                run.bold = True
            else:
                run.text = part
            run.font.color.rgb = RGBColor(39, 43, 103)


def add_macro_to_doc(docx_path):
    abs_path = os.path.abspath(docx_path)

    word = win32.gencache.EnsureDispatch('Word.Application')
    word.Visible = False

    doc = word.Documents.Open(abs_path)

    vb_module = doc.VBProject.VBComponents.Add(1)
    vb_module.CodeModule.AddFromString('''
Sub AutoOpen()
    Dim toc As TableOfContents
    For Each toc In ActiveDocument.TablesOfContents
        toc.Update
    Next toc

    Dim inlineShp As InlineShape
    For Each inlineShp In ActiveDocument.InlineShapes
        inlineShp.ConvertToShape
    Next inlineShp

    Dim shp As Shape
    For Each shp In ActiveDocument.Shapes
        With shp
            .WrapFormat.Type = wdWrapSquare
            .Left = wdShapeLeft
            .RelativeHorizontalPosition = wdRelativeHorizontalPositionMargin
            .Top = wdShapeTop
            .RelativeVerticalPosition = wdRelativeVerticalPositionParagraph
            .WrapFormat.DistanceRight = CentimetersToPoints(0.5)
            .WrapFormat.DistanceBottom = CentimetersToPoints(0.5)
            .WrapFormat.Side = wdWrapRight
            .WrapFormat.AllowOverlap = False
        End With
    Next shp
End Sub
''')

    macro_path = abs_path.replace(".docx", ".docm")
    doc.SaveAs(macro_path, FileFormat=13)

    doc.Close()
    word.Quit()

    return macro_path


def extract_data_by_path(data_dict, path_list):
    try:
        for key in path_list:
            data_dict = data_dict[key]
        return data_dict
    except Exception as e:
        print(f"[extract_data_by_path] Ошибка: {e}")
        return "Нет данных"


async def insert_section(doc, title, level, prompt, source_path=None, page_break=False, subsections=None):
    if page_break:
        doc.add_page_break()

    city = json_data.get("Общие характеристики", {}).get("Город", "Город")
    current_date = datetime.now().strftime("%d.%m.%Y")
    current_year = datetime.now().year

    title_filled = title.format(
        city=city, current_date=current_date, current_year=current_year)

    if level == 1:
        format_text(doc, f"# {title_filled}")
    elif level == 2:
        format_text(doc, f"## {title_filled}")

    if not subsections:
        if prompt.strip():
            data = extract_data_by_path(
                json_data, source_path) if source_path else ""

            if isinstance(data, dict) and "img_src" in data:
                doc.add_paragraph()
                insert_image(doc, data["img_src"])

            prompt_filled = prompt.format(
                city=city,
                current_date=current_date,
                current_year=current_year,
                data=data
            )

            text = await generate_text_gemini(prompt_filled)
            for line in text.split("\n"):
                format_text(doc, line)
        else:
            print(f"Skipping text generation for section: {title_filled}")


def decode_image_base64(base64_str):
    try:
        if ',' in base64_str:
            base64_data = base64_str.split(",")[1]
        else:
            base64_data = base64_str

        base64_data = base64_data.strip().replace('\n', '').replace('\r', '')

        missing_padding = len(base64_data) % 4
        if missing_padding:
            base64_data += '=' * (4 - missing_padding)

        image_data = base64.b64decode(base64_data)

        image = Image.open(BytesIO(image_data))
        image.load()

        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as temp_file:
            image.save(temp_file, format="PNG")
            return temp_file.name

    except Exception as e:
        print(f"[decode_image_base64] Ошибка: {e}")
        raise ValueError("Ошибка при декодировании изображения из base64")


def insert_image(doc, img_src):
    try:
        image_path = decode_image_base64(img_src)

        with Image.open(image_path) as img:
            max_width_in = 4
            max_height_in = 5

            img_width_px, img_height_px = img.size
            aspect_ratio = img_width_px / img_height_px

            width_in = max_width_in
            height_in = max_width_in / aspect_ratio

            if height_in > max_height_in:
                height_in = max_height_in
                width_in = max_height_in * aspect_ratio

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.right_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        run = paragraph.add_run()
        run.add_picture(image_path, width=Inches(
            width_in), height=Inches(height_in))

        os.remove(image_path)

    except Exception as e:
        print(f"Ошибка вставки изображения: {e}")


async def generate_text_gemini(prompt: str) -> str:
    global last_request_time
    try:
        async with semaphore:
            now = time.monotonic()
            wait_time = MIN_INTERVAL - (now - last_request_time)
            if wait_time > 0:
                await asyncio.sleep(wait_time)
            last_request_time = time.monotonic()

            model = genai.GenerativeModel("gemini-2.0-flash")
            response = model.generate_content(prompt)
            return response.text.strip() if response.text else "Ошибка генерации текста"
    except Exception as e:
        return f"Ошибка при генерации текста: {str(e)}"


@app.post("/generate_report")
async def generate_report(request: Request):
    global json_data
    try:
        json_data = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="Ошибка при чтении JSON.")

    report_structure = load_json_file("report_structure.json")

    city = json_data.get("Общие характеристики", {}).get("Город")
    if not city:
        raise HTTPException(
            status_code=500, detail="Город не найден в JSON-файле.")

    doc = Document()
    set_document_styles(doc)
    add_title_page(doc)
    add_table_of_contents(doc)
    add_page_numbers(doc)

    for section in report_structure:
        if "subsections" not in section:
            await insert_section(
                doc,
                title=section["title"],
                level=1,
                prompt=section.get("prompt", ""),
                source_path=section.get("source"),
                page_break=True
            )
        else:
            await insert_section(
                doc,
                title=section["title"],
                level=1,
                prompt="",
                source_path=None,
                page_break=True
            )

            for sub in section["subsections"]:
                await insert_section(
                    doc,
                    title=sub["title"],
                    level=2,
                    prompt=sub.get("prompt", ""),
                    source_path=sub.get("source"),
                    page_break=False
                )

    report_path = os.path.join(REPORTS_DIR, f"Отчет_{city}.docx")
    doc.save(report_path)
    macro_report_path = add_macro_to_doc(report_path)

    return {"report_path": macro_report_path, "city": city}


@app.get("/download_report/")
async def download_report(city: str):
    report_path = os.path.join(REPORTS_DIR, f"Отчет_{city}.docm")
    if os.path.exists(report_path):
        return FileResponse(report_path, filename=f"Отчет_{city}.docm")
    raise HTTPException(status_code=404, detail="Отчет не найден")


@app.get("/")
async def serve_home():
    return FileResponse("static/index.html")

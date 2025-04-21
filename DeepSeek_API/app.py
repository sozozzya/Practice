from fastapi import FastAPI, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
import requests
import os
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime
import re
import base64
import tempfile
from PIL import Image
from io import BytesIO

app = FastAPI()

app.mount("/static", StaticFiles(directory="static"), name="static")

REPORTS_DIR = "reports"
os.makedirs(REPORTS_DIR, exist_ok=True)

DEEPSEEK_API_KEY = "sk-b240d83d2b5f4c5ab5126e1a3905fdbc"

json_file_path = "input.json"
json_data = {}


def load_json_data():
    global json_data
    if not os.path.exists(json_file_path):
        raise HTTPException(status_code=500, detail="JSON-файл не найден.")

    with open(json_file_path, "r", encoding="utf-8") as f:
        json_data = json.load(f)


def set_document_styles(doc):
    sections = doc.sections
    for section in sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    styles = doc.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 147)
    normal_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    normal_style.paragraph_format.first_line_indent = Cm(1.25)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(0)

    heading1 = styles["Heading 1"]
    heading1.font.name = "Times New Roman"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(30, 30, 147)
    heading1.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    heading1.paragraph_format.space_after = Pt(24)

    heading2 = styles["Heading 2"]
    heading2.font.name = "Times New Roman"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(30, 30, 147)
    heading2.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    heading2.paragraph_format.space_before = Pt(24)
    heading2.paragraph_format.space_after = Pt(12)


def add_page_numbers(doc):
    for i, section in enumerate(doc.sections):
        section.different_first_page_header_footer = True

        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph(
        )
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)

        fldSimple = OxmlElement("w:fldSimple")
        fldSimple.set(qn("w:instr"), "PAGE")

        run = paragraph.add_run()
        run._r.append(fldSimple)

        if i == 0:
            section.start_page_number = 2


def add_title_page(doc):
    for _ in range(14):
        doc.add_paragraph()

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(
        "ИССЛЕДОВАНИЕ РЫНКА В СЕГМЕНТЕ, КОТОРОМУ ПРИНАДЛЕЖИТ ОЦЕНИВАЕМАЯ КВАРТИРА")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"

    for _ in range(14):
        doc.add_paragraph()

    for text in [f"Нижний Новгород", str(datetime.today().year)]:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.name = "Times New Roman"


def add_table_of_contents(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run("Содержание")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.bold = True

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldSimple = OxmlElement("w:fldSimple")
    fldSimple.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    run._r.append(fldSimple)


def format_text(doc, text):
    text = text.strip()
    if not text:
        return

    if text.startswith("# "):
        paragraph = doc.add_paragraph(text[2:], style="Heading 1")

    elif text.startswith("## "):
        paragraph = doc.add_paragraph(text[3:], style="Heading 2")

    elif text.startswith("* "):
        paragraph = doc.add_paragraph(style="ListBullet")
        remaining_text = text[2:]

        bold_parts = re.split(r"(\*\*.*?\*\*)", remaining_text)

        for part in bold_parts:
            run = paragraph.add_run()
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]
                run.bold = True
            else:
                run.text = part
            run.font.color.rgb = RGBColor(30, 30, 147)

    else:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        bold_parts = re.split(r"(\*\*.*?\*\*)", text)

        for part in bold_parts:
            run = paragraph.add_run()
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]
                run.bold = True
            else:
                run.text = part
            run.font.color.rgb = RGBColor(30, 30, 147)


def decode_image_base64(base64_str):
    try:
        if ',' in base64_str:
            base64_data = base64_str.split(",")[1]
        else:
            base64_data = base64_str

        base64_data = base64_data.strip().replace('\n', '').replace('\r', '')

        missing_padding = len(base64_data) % 4
        if missing_padding:
            base64_data += '=' * (4 - missing_padding)

        image_data = base64.b64decode(base64_data)

        image = Image.open(BytesIO(image_data))
        image.load()

        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as temp_file:
            image.save(temp_file, format="PNG")
            return temp_file.name

    except Exception as e:
        print(f"[decode_image_base64] Ошибка: {e}")
        raise ValueError("Ошибка при декодировании изображения из base64")


def insert_image(doc, img_src):
    try:
        image_path = decode_image_base64(img_src)

        with Image.open(image_path) as img:
            max_width = Inches(5)
            max_height = Inches(5)

            img_width, img_height = img.size
            aspect_ratio = img_width / img_height

            width = max_width
            height = max_width / aspect_ratio

            if height > max_height:
                height = max_height
                width = max_height * aspect_ratio

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = paragraph.add_run()
        run.add_picture(image_path, width=width, height=height)

        os.remove(image_path)

    except Exception as e:
        print(f"Ошибка вставки изображения: {e}")


def generate_text_deepseek(prompt: str) -> str:
    try:
        url = "https://api.deepseek.com/v1/chat/completions"

        headers = {
            "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
            "Content-Type": "application/json"
        }

        data = {
            "model": "deepseek-chat",
            "messages": [
                {"role": "user", "content": prompt}
            ],
            "temperature": 0.7
        }

        response = requests.post(url, headers=headers, json=data)
        response.raise_for_status()
        result = response.json()

        return result["choices"][0]["message"]["content"].strip()

    except Exception as e:
        return f"Ошибка при генерации текста: {str(e)}"


@app.post("/generate_report")
async def generate_report():
    load_json_data()

    city = json_data.get("Общие характеристики", {}).get("Город")
    if not city:
        raise HTTPException(
            status_code=500, detail="Город не найден в JSON-файле.")

    doc = Document()
    set_document_styles(doc)
    add_title_page(doc)
    add_table_of_contents(doc)
    add_page_numbers(doc)

    current_date = datetime.now().strftime("%d.%m.%Y")

    report_structure = [
        {
            "title": "1. Анализ влияния общей политической и социально-экономической обстановки в России на рынок продажи и аренды недвижимости, в том числе тенденций, наметившихся на рынке за 2024 год",
            "prompt": f"Составь подробный отчет по анализу влияния общей политической и социально-экономической обстановки в России на рынок недвижимости. Укажи достоверные цифры на {current_date} и приведи конкретные примеры. Для анализа использовать публикации только надежных источников. Укажи ссылки на используемые публикации."
        },
        {
            "title": "2. Анализ влияния общей политической и социально-экономической обстановки в городе Нижний Новгород на рынок продажи и аренды недвижимости, в том числе тенденций, наметившихся на рынке за 2024 год",
            "prompt": f"Составь подробный отчет по анализу влияния общей политической и социально-экономической обстановки в городе {city} на рынок недвижимости. Укажи достоверные цифры на {current_date} и приведи конкретные примеры. Для анализа использовать публикации только надежных источников. Укажи ссылки на используемые публикации."
        },
        {
            "title": "3. Объект оценки",
            "subsections": [
                {
                    "title": "3.1. Общие сведения",
                    "prompt": f"Составь подробный отчет по анализу общих сведений об объекте оценки. Для этого проанализируй данные: {json_data.get('Общие сведения', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Общие сведения', 'Нет данных')
                },
                {
                    "title": "3.2. Общие характеристики",
                    "prompt": f"Составь подробный отчет по анализу общих характеристик объекта оценки. Для этого проанализируй данные: {json_data.get('Общие характеристики', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Общие характеристики', 'Нет данных')
                },
                {
                    "title": "3.3. Характеристики здания",
                    "prompt": f"Составь подробный отчет по анализу общих характеристик здания, в котором находится объект оценки. Для этого проанализируй данные: {json_data.get('Характеристики здания', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Характеристики здания', 'Нет данных')
                }
            ]
        },
        {
            "title": "4. Анализ местоположения объекта оценки",
            "subsections": [
                {
                    "title": "4.1. Территориально функциональная зона",
                    "prompt": f"Составь подробный отчет по анализу территориально-функциональной зоны, в котором находится объект оценки. Для этого проанализируй данные: {json_data.get('Зона', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Зона', 'Нет данных')
                },
                {
                    "title": "4.2. Рейтинг зоны местонахождения",
                    "prompt": f"Составь подробный отчет по анализу рейтинга территориально-функциональной зоны, в котором находится объект оценки. Для этого проанализируй данные: {json_data.get('Рейтинг зоны местонахождения', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Рейтинг зоны местонахождения', 'Нет данных')
                },
                {
                    "title": "4.3. Ближайшее окружение",
                    "prompt": f"Составь подробный отчет по анализу ближайшего окружения объекта оценки. Для этого проанализируй данные: {json_data.get('Анализ местоположения', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ местоположения', 'Нет данных')
                }
            ]
        },
        {
            "title": "5. Анализ фактических данных о ценах, предложений и арендных ставок с объектами из сегмента рынка, к которому может быть оцениваемый объект",
            "subsections": [
                {
                    "title": "5.1. Фактические данные",
                    "prompt": f"Составь подробный отчет по анализу фактических данных о ценах и арендных ставках в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Фактические данные', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Фактические данные', 'Нет данных')
                },
                {
                    "title": "5.2. Удельная цена",
                    "prompt": f"Составь подробный отчет по анализу удельной цены квартир в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Удельная цена', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Удельная цена', 'Нет данных')
                },
                {
                    "title": "5.3. Удельная арендная ставка",
                    "prompt": f"Составь подробный отчет по анализу удельной арендной ставки на квартиры в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Удельная арендная ставка', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Удельная арендная ставка', 'Нет данных')
                },
                {
                    "title": "5.4. Общая площадь",
                    "prompt": f"Составь подробный отчет по анализу общей площади квартир в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Общая площадь', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Общая площадь', 'Нет данных')
                },
                {
                    "title": "5.5. Состояние отделки",
                    "prompt": f"Составь подробный отчет по анализу состояния отделки квартир в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Состояние отделки', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Состояние отделки', 'Нет данных')
                },
                {
                    "title": "5.6. Количество комнат",
                    "prompt": f"Составь подробный отчет по анализу количества комнат в квартирах в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Количество комнат', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Количество комнат', 'Нет данных')
                },
                {
                    "title": "5.7. Количество просмотров",
                    "prompt": f"Составь подробный отчет по анализу количества просмотров на квартиры в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ рынка', {}).get('Количество просмотров', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ рынка', {}).get('Количество просмотров', 'Нет данных')
                }
            ]
        },
        {
            "title": "6. Динамика рынка",
            "subsections": [
                {
                    "title": "6.1. Усредненные показатели за год",
                    "prompt": f"Составь подробный отчет по анализу усредненных показателей за год динамики рынка квартир в городе {city}. Для этого проанализируй данные: {json_data.get('Динамика рынка', {}).get('Усредненные показатели за год', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Динамика рынка', {}).get('Усредненные показатели за год', 'Нет данных')
                },
                {
                    "title": "6.2. Динамика средней удельной цены",
                    "prompt": f"Составь подробный отчет по анализу динамики средней удельной цены квартир в городе {city}. Для этого проанализируй данные: {json_data.get('Динамика рынка', {}).get('Удельная цена', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Динамика рынка', {}).get('Удельная цена', 'Нет данных')
                },
                {
                    "title": "6.3. Динамика средней удельной арендной ставки",
                    "prompt": f"Составь подробный отчет по анализу динамики средней удельной арендной ставки на квартиры в городе {city}. Для этого проанализируй данные: {json_data.get('Динамика рынка', {}).get('Удельная арендная ставка', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Динамика рынка', {}).get('Удельная арендная ставка', 'Нет данных')
                },
                {
                    "title": "6.4. Динамика валового мультипликатора",
                    "prompt": f"Составь подробный отчет по анализу динамики валового мультипликатора в городе {city}. Для этого проанализируй данные: {json_data.get('Динамика рынка', {}).get('Валовый мультипликатор', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Динамика рынка', {}).get('Валовый мультипликатор', 'Нет данных')
                },
                {
                    "title": "6.5. Динамика доходности",
                    "prompt": f"Составь подробный отчет по анализу доходности в городе {city}. Для этого проанализируй данные: {json_data.get('Динамика рынка', {}).get('Доходность', 'Нет данных')}. Опиши их связным текстом».",
                    "source": json_data.get('Динамика рынка', {}).get('Доходность', 'Нет данных')
                }
            ]
        },
        {
            "title": "7. Анализ основных факторов, влияющих на цены и (или) арендные ставки сопоставимых с оцениваемым объектов недвижимости",
            "prompt": f"Составь подробный отчет по анализу основных факторов, влияющих на цены недвижимости в городе {city}. Укажи достоверные цифры на {current_date} и приведи конкретные примеры. Для анализа использовать публикации только надежных источников. Укажи ссылки на используемые публикации."
        },
        {
            "title": "8. Анализ активности продавцов",
            "subsections": [
                {
                    "title": "8.1. Количество объектов, выставленных на продажу и актуальных на начало наблюдений",
                    "prompt": f"Составь подробный отчет по анализу количества объектов, выставленных на продажу и актуальных на начало наблюдений, в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('n1_n2', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('n1_n2', 'Нет данных')
                },
                {
                    "title": "8.2. Количество объектов, выставленных на продажу и вновь появившихся в периоде",
                    "prompt": f"Составь подробный отчет по анализу количества объектов, выставленных на продажу и и вновь появившихся в периоде, в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('n3_n4', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('n3_n4', 'Нет данных')
                },
                {
                    "title": "8.3. Активность продавцов в периоде",
                    "prompt": f"Составь подробный отчет по анализу активности продавцов в периоде (отношение новых объявлений к актуальным на начала периода (n/N₀)) в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('a', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('a', 'Нет данных')
                },
                {
                    "title": "8.4. Активность покупателей в периоде",
                    "prompt": f"Составь подробный отчет по анализу активности покупателей в периоде (отношение снятых объявлений к актуальным на начало периода(m/N₀)) в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('m_', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('m_', 'Нет данных')
                },
                {
                    "title": "8.5. Среднее количество просмотров у данных, выставленных на продажу и вновь появившихся",
                    "prompt": f"Составь подробный отчет по анализу среднего количества просмотров у данных, выставленных на продажу и вновь появившихся в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('p3_p4_mean', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('p3_p4_mean', 'Нет данных')
                },
                {
                    "title": "8.6. Интенсивность продаж данных, выставленных на продажу и актуальных на начало наблюдений",
                    "prompt": f"Составь подробный отчет по анализу интенсивности продаж данных, выставленных на продажу и актуальных на начало наблюдений в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('t1', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('t1', 'Нет данных')
                },
                {
                    "title": "8.7. Интенсивность продаж данных, выставленных на продажу и вновь появившихся",
                    "prompt": f"Составь подробный отчет по анализу интенсивности продаж данных, выставленных на продажу и вновь появившихся в городе {city}. Для этого проанализируй данные: {json_data.get('Анализ активности продавцов', {}).get('t2', 'Нет данных')}. Опиши их связным текстом.",
                    "source": json_data.get('Анализ активности продавцов', {}).get('t2', 'Нет данных')
                }
            ]
        },
        {
            "title": "9. INF-оценкка",
            "prompt": f"«Составь подробный отчет об оценке квартиры. Для этого проанализируй данные: {json_data.get('INF-Оценка', 'Нет данных')}. Опиши их связным текстом.",
            "source": json_data.get('INF-Оценка', 'Нет данных')
        }
    ]

    for section in report_structure:
        if "prompt" in section:
            doc.add_page_break()
            format_text(doc, f"# {section['title']}")

            source = section.get("source", {})
            if isinstance(source, dict) and "img_src" in source:
                insert_image(doc, source["img_src"])

            prompt = section["prompt"]
            text = generate_text_deepseek(prompt)
            for line in text.split("\n"):
                format_text(doc, line)

        if "subsections" in section:
            doc.add_page_break()
            format_text(doc, f"# {section['title']}")
            for sub in section["subsections"]:
                format_text(doc, f"## {sub['title']}")

                source = sub.get("source", {})
                if isinstance(source, dict) and "img_src" in source:
                    insert_image(doc, source["img_src"])

                prompt = sub["prompt"]
                text = generate_text_deepseek(prompt)
                for line in text.split("\n"):
                    format_text(doc, line)

    report_path = os.path.join(REPORTS_DIR, f"report_{city}.docx")
    doc.save(report_path)
    return {"report_path": report_path, "city": city}


@app.get("/download_report/")
async def download_report(city: str):
    report_path = os.path.join(REPORTS_DIR, f"report_{city}.docx")
    if os.path.exists(report_path):
        return FileResponse(report_path, filename=f"report_{city}.docx")
    raise HTTPException(status_code=404, detail="Отчет не найден")


@app.get("/")
async def serve_home():
    return FileResponse("static/index.html")
